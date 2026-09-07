"""
Exportación a Excel con la hoja de procedencia de los datos.

Un archivo de datos sin la consulta que lo produjo no se puede reproducir ni
auditar: a los tres meses nadie recuerda qué filtros se aplicaron. Por eso el
libro lleva siempre dos hojas, datos y consulta.

Monitoría de investigación - Beca Avanza, Universidad de los Andes.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

FUENTE = "Arial"
FORMATO_MONEDA = "#,##0"
ANCHO_MAXIMO = 42


def _describir_filtros(filtros: dict | None) -> list[tuple[str, str]]:
    """Convierte el diccionario de filtros en filas legibles."""
    filas: list[tuple[str, str]] = []
    for columna, valor in (filtros or {}).items():
        if isinstance(valor, dict):
            desde, hasta = valor.get("desde", "—"), valor.get("hasta", "—")
            filas.append((f"Filtro · {columna}", f"de {desde} a {hasta}"))
        else:
            filas.append((f"Filtro · {columna}", str(valor)))
    return filas or [("Filtros", "ninguno (se trajo todo lo disponible)")]


def hoja_consulta(metadatos: dict) -> pd.DataFrame:
    """
    Arma la hoja que documenta de dónde salieron los datos.

    metadatos admite: tabla, dataset, modo, filtros, limite, consulta,
    filas, columnas y margen_meses.
    """
    filas: list[tuple[str, str]] = [
        ("Fecha de extracción", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Fuente", "SECOP 2 · datos.gov.co (Socrata Open Data API)"),
        ("Tabla", str(metadatos.get("tabla", "—"))),
        ("Identificador del dataset", str(metadatos.get("dataset", "—"))),
        ("Modo", str(metadatos.get("modo", "—"))),
    ]
    filas += _describir_filtros(metadatos.get("filtros"))

    limite = metadatos.get("limite")
    filas += [
        ("Tope de filas", "sin tope" if limite in (None, "") else f"{limite:,}".replace(",", ".")),
        ("Filas obtenidas", f"{metadatos.get('filas', 0):,}".replace(",", ".")),
        ("Columnas", str(metadatos.get("columnas", "—"))),
    ]
    if metadatos.get("margen_meses") is not None:
        filas.append(("Margen para buscar procesos",
                      f"{metadatos['margen_meses']} meses antes"))
    if metadatos.get("consulta"):
        filas.append(("Consulta enviada", str(metadatos["consulta"])))

    filas.append((
        "Advertencia",
        "Los indicadores de riesgo son señales de priorización, no evidencia "
        "de irregularidad. Varios son el comportamiento normal de ciertas "
        "modalidades de contratación.",
    ))
    return pd.DataFrame(filas, columns=["Campo", "Valor"])


def libro_excel(
    df: pd.DataFrame,
    metadatos: dict,
    columnas_moneda: list[str] | None = None,
) -> bytes:
    """Devuelve el archivo .xlsx con las hojas «Datos» y «Consulta»."""
    buffer = io.BytesIO()
    columnas_moneda = columnas_moneda or []

    with pd.ExcelWriter(buffer, engine="openpyxl") as escritor:
        df.to_excel(escritor, sheet_name="Datos", index=False)
        hoja_consulta(metadatos).to_excel(escritor, sheet_name="Consulta", index=False)

        libro = escritor.book
        datos, consulta = libro["Datos"], libro["Consulta"]

        for hoja in (datos, consulta):
            for celda in hoja[1]:
                celda.font = Font(name=FUENTE, bold=True)
            hoja.freeze_panes = "A2"

        # Ancho por el encabezado, que es lo que se necesita leer; el contenido
        # se ve al ensanchar la columna si hace falta.
        for indice, nombre in enumerate(df.columns, start=1):
            datos.column_dimensions[get_column_letter(indice)].width = min(
                max(12, len(str(nombre)) + 2), ANCHO_MAXIMO
            )
            if nombre in columnas_moneda:
                for celda in datos[get_column_letter(indice)][1:]:
                    celda.number_format = FORMATO_MONEDA

        consulta.column_dimensions["A"].width = 30
        consulta.column_dimensions["B"].width = 90
        for fila in consulta.iter_rows(min_row=2):
            fila[0].font = Font(name=FUENTE, bold=True)
            fila[1].alignment = Alignment(wrap_text=True, vertical="top")

    logger.info("Libro de Excel generado: %d filas, %d columnas", len(df), df.shape[1])
    return buffer.getvalue()
